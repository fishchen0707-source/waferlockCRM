# 把「下單規格表」資料夾裡 106 款鎖的編碼參數表，解析成一份機器可讀的料號參數字典。
#
# ── 為什麼需要這支 ────────────────────────────────────────
# 料號不是查出來的，是**組**出來的：
#   L396-1C17C1-A0122C-A1CA1-2X30A
#   └鎖 └型號 └前飾板/前面板/… └鎖腹方向/鎖仁/門厚 └…
# 每一段是一個參數，每個參數有自己的代碼表。光 C762 一款理論組合就約 100 萬種，
# 所以「從 ERP 匯出一份料號清單」不可能涵蓋——必須拿到**編碼規則**本身。
#
# ── 為什麼用「遮罩列」當錨點，而不是找「成品代碼」四個字 ──────
# 第一版是找「成品代碼」，結果 31 款抓不到，而且幾乎全是 L 系列
# （L 系列的欄位名是「鎖／產品型號」，不是「成品代碼／機構碼」）。
# 真正跨系列一致的是**遮罩列**——那一列長得像 `L | 396 | — | X | X | XX`，
# 整列都是單一字元、X 串或短數字，參數名永遠在它的下一列。
# 改用這個規則之後 106/106 全過。
#
#   python scripts/parse-order-specs.py
import collections
import json
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")          # openpyxl 對資料驗證的警告，與解析無關
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIRS = [os.path.join(ROOT, "下單規格表"), os.path.join(ROOT, "下單規格表-xlsx轉檔")]
OUT = os.path.join(ROOT, "docs", "料號參數字典.json")

FNAME = re.compile(r"^(W3-MS001-M-\d+)\s+(.+?)下單規格表\s*V?([\d.]+)\.xlsx$", re.I)
DASH = {"—", "-", "–", "─"}
SCAN_ROWS = 20                              # 遮罩列一定在前 20 列內（實測最深第 4 列）

# 「代碼」與「說明」之間的分隔符不統一：多數用冒號，L320 系列用底線、L390 用點。
# 不能只認冒號——第一版只認冒號時，有 49 款出現「拆不開」的選項。
# 代碼最長放到 10：C720 的「類型/防鑽/長度/顏色」合成一段，代碼是 S03030A（7 碼）。
SEP = re.compile(r"^([A-Za-z0-9]{1,10})\s*[:：_.]\s*(.+)$")

# 規格表的選項欄裡混著**註解與異動紀錄**，那些不是可選項目：
#   「**標準為IP54」「**SBA與電梯控制器共用PCB及韌體」
#   「24/5/15 以下細天線設隱藏, 新增粗天線」（日期開頭的修訂紀錄）
# 混進下拉選單業務會看不懂，所以在這裡濾掉——但只濾這兩種明確的樣式，
# 其餘拆不開的仍然照收並回報，不擴大解釋。
NOTE = re.compile(r"^(?:[*＊]|\d{1,2}/\d{1,2}/\d{1,2}\s)")

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def is_number_row(r):
    """段位編號列：整列都是純數字（L310 在參數名下面多一列 000/010/020…）。
    這種列不是選項，混進去會讓下拉選單多出一堆看不懂的數字。"""
    nz = [("" if c is None else str(c).strip()) for c in r]
    nz = [v for v in nz if v]
    return len(nz) >= 5 and all(re.fullmatch(r"\d{2,4}", v) for v in nz)


def is_mask_cell(v):
    """遮罩列的儲存格長相：單一字元、X 串、或 2~5 位數字，或分隔線。"""
    return v in DASH or bool(re.fullmatch(r"[A-Za-z0-9]|[Xx]{1,6}|\d{2,5}", v))


def latest_files():
    """同一款有多版時取版號最大的那一份。"""
    best = {}
    for d in DIRS:
        if not os.path.isdir(d):
            continue
        for f in os.listdir(d):
            if f.startswith("~$"):          # Excel 開啟中的暫存檔
                continue
            m = FNAME.match(f)
            if not m:
                continue
            model, ver = m.group(2), m.group(3)
            vt = tuple(int(x) for x in ver.split("."))
            if model not in best or vt > best[model][0]:
                best[model] = (vt, m.group(1), ver, os.path.join(d, f))
    return best


def parse(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = next((s for s in wb.sheetnames if "編碼參數" in s), None)
    if not sheet:
        wb.close()
        return None, "找不到「編碼參數表」分頁"
    ws = wb[sheet]
    rows = [[("" if c is None else str(c).strip()) for c in r]
            for r in ws.iter_rows(values_only=True)]
    wb.close()

    mask_i = None
    for i, vals in enumerate(rows[:SCAN_ROWS]):
        nz = [v for v in vals if v]
        if len(nz) >= 5 and sum(1 for v in nz if is_mask_cell(v)) >= len(nz) * 0.8:
            mask_i = i
            break
    if mask_i is None:
        return None, "找不到遮罩列"
    if mask_i + 1 >= len(rows):
        return None, "遮罩列之後沒有參數名列"

    mask, names = rows[mask_i], rows[mask_i + 1]
    segs, issues = [], []
    for col in range(len(mask)):
        mk = mask[col] if col < len(mask) else ""
        if not mk:
            continue
        nm = names[col] if col < len(names) else ""
        if mk in DASH:
            segs.append({"index": col, "mask": mk, "name": "", "type": "sep", "options": []})
            continue

        opts = []
        for r in rows[mask_i + 2:]:
            if is_number_row(r):
                continue                      # 段位編號列（000/010/020…），不是選項
            v = r[col] if col < len(r) else ""
            if not v or v == "`":
                continue
            if NOTE.match(v):
                continue
            m2 = SEP.match(v)
            if m2:
                opts.append({"code": m2.group(1).strip(), "desc": m2.group(2).strip()})
            else:
                # 拆不開的照收並標記，不自己猜——猜錯會讓業務選到錯的料號
                opts.append({"code": v, "desc": ""})
                issues.append(f"欄{col}「{nm or mk}」的選項『{v[:24]}』拆不出代碼與說明")
        # 遮罩是字面值（L、396、762）＝這一段是固定的，不給人選。
        # 遮罩是 X 串＝要選。分清楚，產生器才知道哪幾段要顯示成下拉。
        literal = not re.fullmatch(r"[Xx]{1,6}", mk)
        segs.append({"index": col, "mask": mk, "name": nm,
                     "type": "fixed" if literal and len(opts) <= 1 else "choice",
                     "options": opts})
    return {"mask_row": mask_i + 1, "segments": segs}, issues


def main():
    files = latest_files()
    print(f"找到 {len(files)} 款（同款多版取最新）")
    out, failed, all_issues = {}, [], collections.Counter()

    for model, (vt, no, ver, path) in sorted(files.items()):
        data, issues = parse(path)
        if data is None:
            failed.append(f"{model}：{issues}")
            continue
        combos = 1
        for s in data["segments"]:
            if s["options"]:
                combos *= len(s["options"])
        out[model] = {
            "doc_no": no, "version": ver,
            "file": os.path.relpath(path, ROOT).replace("\\", "/"),
            "mask_row": data["mask_row"],
            "segment_count": len(data["segments"]),
            "combinations": combos,
            "segments": data["segments"],
        }
        for i in (issues or []):
            all_issues[model] += 1

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    print(f"解析成功 {len(out)} 款 → {os.path.relpath(OUT, ROOT)}"
          f"（{os.path.getsize(OUT) / 1024:.0f} KB）")
    if failed:
        print(f"\n🔴 失敗 {len(failed)} 款：")
        for x in failed:
            print("  " + x)
    if all_issues:
        print(f"\n⚠ 有選項不符「代碼:說明」格式的 {len(all_issues)} 款"
              f"（已照收並標記，沒有自己猜怎麼拆）：")
        for m, n in all_issues.most_common(12):
            print(f"  {m}：{n} 個")


if __name__ == "__main__":
    main()
